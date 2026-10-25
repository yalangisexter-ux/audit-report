"""Store Directory regression tests (restored after truncated push)."""
import hashlib

import openpyxl

import customer_prospects as prospects


def _make_reference_workbooks(tmp_path):
    database_path = tmp_path / "customer database.xlsx"
    database = openpyxl.Workbook()
    stores = database.active
    stores.title = "Haryana stores"
    stores.append([
        "Business Name", "Business Category", "State", "City", "Mobile Number",
        "Connected/Not Connected", "Verified (Yes/No)", "Google Maps Link",
    ])
    stores.append(["North Fitness Hub", "Sports retailer", "Haryana", "Gurugram", "9000000001", "Connected", "Yes", "https://maps.example/store"])
    stores.append(["Business Name", "Business Category", "State", "City", "Mobile Number", "Connected/Not Connected", "Verified (Yes/No)", "Google Maps Link"])
    stores.append([None, "section note", None, None, None, None, None, None])

    gyms = database.create_sheet("Punjab gym")
    gyms.append(["Business Name", "Business Category", "Contact Person", "Phone Number", "State", "City", "Address", "Google Maps"])
    gyms.append(["Bodyweight Studio", "Gym", "Manager", "9000000002", "Punjab", "Ludhiana", "Main Road", "https://maps.example/gym"])
    database.save(database_path)

    dashboard_path = tmp_path / "store dashboard.xlsx"
    dashboard = openpyxl.Workbook()
    sheet = dashboard.active
    sheet.append(["State/UT", "Stores"])
    sheet.append(["Haryana", 2])
    sheet.append(["Punjab", 1])
    sheet.append(["TOTAL", 3])
    dashboard.save(dashboard_path)
    return database_path, dashboard_path


def _sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_workbooks_load_read_only_and_normalize_regional_schemas(tmp_path):
    database_path, dashboard_path = _make_reference_workbooks(tmp_path)
    database_before = _sha256(database_path)
    dashboard_before = _sha256(dashboard_path)

    records, stats = prospects.parse_database_workbook(database_path)
    dashboard_counts, dashboard_total = prospects.parse_dashboard_workbook(dashboard_path)

    assert len(records) == 2
    assert records[0]["kind"] == "Store"
    assert records[0]["business_name"] == "North Fitness Hub"
    assert records[0]["phone"] == "9000000001"
    assert records[1]["kind"] == "Gym"
    assert records[1]["state"] == "Punjab"
    assert stats == {"scanned_rows": 3, "skipped_unidentified": 1}
    assert dashboard_counts == {"Haryana": 2, "Punjab": 1}
    assert dashboard_total == 3
    assert _sha256(database_path) == database_before
    assert _sha256(dashboard_path) == dashboard_before


def test_negative_filter_excludes_pharmacy_herbal_healthkart():
    assert prospects.NEGATIVE_FILTER_PATTERN.search("Apollo Pharmacy Store")
    assert prospects.NEGATIVE_FILTER_PATTERN.search("Herbalife Nutrition Center")
    assert prospects.NEGATIVE_FILTER_PATTERN.search("HealthKart Store Rohini")
    assert not prospects.NEGATIVE_FILTER_PATTERN.search("Alpha Sports Nutrition")
    assert not prospects.NEGATIVE_FILTER_PATTERN.search("Fit Fuel Supplements")


def test_customer_prospects_filter_options_for_stores_gyms_and_hybrids(tmp_path, monkeypatch):
    monkeypatch.setenv("QT_QPA_PLATFORM", "offscreen")
    from PySide6.QtWidgets import QApplication

    qt_app = QApplication.instance() or QApplication([])
    tab = prospects.CustomerProspectsTab()
    try:
        tab.kind_filter.setCurrentText("All types")
        # Pure Product Selling Stores filter is strict: product evidence + full
        # postal address (with PIN) + actionable Maps link are required.
        tab.records = [
            {
                "kind": "Store",
                "business_name": "Alpha Nutrition Store",
                "state": "Delhi",
                "city": "New Delhi",
                "business_category": "Sports Nutrition / Supplement Store",
                "address": "12 Main Market Road, Connaught Place, New Delhi 110001",
                "maps": "https://www.google.com/maps/search/?api=1&query=Alpha+Nutrition+Store",
                "phone": "9999999999",
            },
            {
                "kind": "Gym",
                "business_name": "Titan Gym & Fitness",
                "state": "Delhi",
                "city": "New Delhi",
                "business_category": "Fitness Centre / Gym",
                "address": "45 Gym Street, South Delhi 110017",
                "maps": "https://maps.google.com/?q=Titan+Gym",
                "phone": "8888888888",
            },
            {
                "kind": "Gym + Store",
                "business_name": "Iron Beast Gym Pro Shop",
                "state": "Delhi",
                "city": "New Delhi",
                "business_category": "Gym with Nutrition & Supplement Counter / Pro-Shop",
                "address": "78 Hybrid Lane, East Delhi 110092",
                "maps": "https://maps.google.com/?q=Iron+Beast",
                "phone": "7777777777",
            },
        ]
        tab.contact_filter.setCurrentText("All contacts")
        tab._fill_leads_table()

        tab.kind_filter.setCurrentText("All types")
        tab.apply_filters()
        assert sum(not tab.leads_table.isRowHidden(i) for i in range(3)) == 3

        tab.kind_filter.setCurrentText("Pure Product Selling Stores")
        tab.apply_filters()
        assert not tab.leads_table.isRowHidden(0)
        assert tab.leads_table.isRowHidden(1)
        assert tab.leads_table.isRowHidden(2)

        tab.kind_filter.setCurrentText("Gym")
        tab.apply_filters()
        assert tab.leads_table.isRowHidden(0)
        assert not tab.leads_table.isRowHidden(1)
        assert tab.leads_table.isRowHidden(2)

        tab.kind_filter.setCurrentText("Gyms which sell products also")
        tab.apply_filters()
        assert tab.leads_table.isRowHidden(0)
        assert tab.leads_table.isRowHidden(1)
        assert not tab.leads_table.isRowHidden(2)
    finally:
        tab.close()


def test_deduplication_and_enrichment_prevents_duplicate_entries(tmp_path, monkeypatch):
    import app

    monkeypatch.setattr(app, "DB_PATH", str(tmp_path / "dedupe_test.db"))
    monkeypatch.setattr(app, "BACKUP_DIR", str(tmp_path / "backups"))
    db = app.Database()
    initial_count = len(prospects.load_saved_prospects(db))

    initial_records = [
        {
            "kind": "Store",
            "business_name": "Fitnessstack: India&#x27;s Nutrition Store",
            "state": "Karnataka",
            "city": "Bengaluru",
            "phone": "",
            "email": "",
            "address": "Indiranagar, Bangalore",
            "website": "https://fitnessstack.com",
            "maps": "",
            "notes": "User custom lead note - do not overwrite",
        }
    ]
    res1 = prospects.persist_prospect_records(db, initial_records)
    assert res1["inserted"] == 1
    assert res1["duplicates"] == 0

    saved = prospects.load_saved_prospects(db)
    assert len(saved) == initial_count + 1
    matching = [s for s in saved if s["business_name"] == "Fitnessstack: India's Nutrition Store"]
    assert len(matching) == 1
    assert matching[0]["phone"] == ""

    incoming_crawl = [
        {
            "kind": "Store",
            "business_name": "Fitnessstack: India's Nutrition Store",
            "state": "Karnataka",
            "city": "Bengaluru",
            "phone": "+91 98765 43210",
            "email": "contact@fitnessstack.com",
            "address": "100 Feet Road, Indiranagar, Bengaluru, Karnataka 560038",
            "website": "https://www.fitnessstack.com/shop",
            "maps": "https://maps.google.com/?q=fitnessstack",
            "notes": "Scraper raw notes",
        }
    ]
    res2 = prospects.persist_prospect_records(db, incoming_crawl)
    assert res2["inserted"] == 0
    assert res2["duplicates"] == 1

    saved_after = prospects.load_saved_prospects(db)
    assert len(saved_after) == initial_count + 1

    matching_after = [s for s in saved_after if s["business_name"] == "Fitnessstack: India's Nutrition Store"]
    assert len(matching_after) == 1
    enriched = matching_after[0]
    assert enriched["phone"] == "+91 98765 43210"
    assert enriched["email"] == "contact@fitnessstack.com"
    assert "100 Feet Road" in enriched["address"]
    # Existing Google Maps URLs are kept; only blank/invalid maps are rewritten.
    assert enriched["maps"].startswith((
        "https://www.google.com/maps/",
        "https://maps.google.com/",
        "https://maps.app.goo.gl/",
    ))
    assert "fitnessstack" in enriched["maps"].casefold()
    assert enriched["notes"] == "User custom lead note - do not overwrite"


def test_territory_separation_states_uts_islands():
    assert len(prospects.INDIA_STATES) == 28
    assert len(prospects.INDIA_UNION_TERRITORIES) == 6
    assert len(prospects.INDIA_ISLAND_TERRITORIES) == 2
    assert len(prospects.INDIA_STATES) + len(prospects.INDIA_UNION_TERRITORIES) + len(prospects.INDIA_ISLAND_TERRITORIES) == 36
    assert prospects.territory_type("Haryana") == "State"
    assert prospects.territory_type("Delhi") == "Union Territory"
    assert prospects.territory_type("Andaman and Nicobar Islands") == "Island Territory"
