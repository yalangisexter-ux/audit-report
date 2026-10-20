import csv
from datetime import date, datetime

import app


def _make_window(tmp_path, monkeypatch):
    monkeypatch.setenv("QT_QPA_PLATFORM", "offscreen")
    qt_app = app.QApplication.instance() or app.QApplication([])
    monkeypatch.setattr(app, "DB_PATH", str(tmp_path / "ui-print.db"))
    monkeypatch.setattr(app, "BACKUP_DIR", str(tmp_path / "backups"))
    database = app.Database()
    window = app.MainWindow(database, {"username": "Print QA", "role": "Admin"})
    return qt_app, database, window


def _insert_business_data(database):
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    database.conn.execute(
        """INSERT INTO bt_products(name,sku,category,brand,base_unit,opening_stock,
           received_qty,reserved_qty,reorder_level,selling_price,updated_at)
           VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
        ("QA Protein Bar", "QA-PB-1", "Nutrition", "Raw Strength", "PCS", 10, 5, 1, 4, 89.5, timestamp),
    )
    database.conn.execute(
        """INSERT INTO bt_services(name,category,description,training_mode,service_unit,
           session_duration,num_sessions,plan_validity,standard_price,active,updated_at)
           VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
        ("QA Calisthenics Plan", "Training", "Four-week starter plan", "Hybrid", "Session", 45, 12, 28, 2400, 1, timestamp),
    )
    database.conn.commit()


def test_product_and_service_print_checkboxes_filter_all_export_formats(tmp_path, monkeypatch):
    from docx import Document
    import reportlab.platypus

    qt_app, database, window = _make_window(tmp_path, monkeypatch)
    try:
        _insert_business_data(database)
        today = date.today().isoformat()
        options = window.selected_eod_options()
        assert options["business_products"] is False
        assert options["business_services"] is False
        assert not window.business_data_group.isHidden()

        window.business_report_checks["business_products"].setChecked(True)
        window.business_report_checks["business_services"].setChecked(True)
        options = window.selected_eod_options()
        assert options["business_products"] is True
        assert options["business_services"] is True

        docx_path = tmp_path / "with-business-data.docx"
        csv_path = tmp_path / "with-business-data.csv"
        pdf_path = tmp_path / "with-business-data.pdf"
        window.make_docx(docx_path, today, report_type="Employee Details", eod_options=options)
        window.export_report_csv(csv_path, today, report_type="Employee Details", eod_options=options)
        window.make_pdf(pdf_path, today, report_type="Employee Details", eod_options=options)

        word = Document(docx_path)
        word_text = "\n".join(
            [paragraph.text for paragraph in word.paragraphs]
            + [cell.text for table in word.tables for row in table.rows for cell in row.cells]
        )
        assert "QA Protein Bar" in word_text
        assert "QA Calisthenics Plan" in word_text
        assert "14" in word_text  # available stock: 10 + 5 - 1

        with csv_path.open(encoding="utf-8-sig", newline="") as stream:
            csv_text = "\n".join(",".join(row) for row in csv.reader(stream))
        assert "QA Protein Bar" in csv_text
        assert "QA Calisthenics Plan" in csv_text
        assert pdf_path.is_file() and pdf_path.stat().st_size > 1000

        captured = []

        class CaptureDocument:
            def __init__(self, *args, **kwargs):
                pass

            def build(self, story, **kwargs):
                captured.extend(story)

        monkeypatch.setattr(reportlab.platypus, "SimpleDocTemplate", CaptureDocument)
        window.make_pdf(tmp_path / "captured.pdf", today, report_type="Employee Details", eod_options=options)
        pdf_story_text = "\n".join(
            flowable.getPlainText() for flowable in captured if hasattr(flowable, "getPlainText")
        )
        assert "QA Protein Bar" in pdf_story_text
        assert "QA Calisthenics Plan" in pdf_story_text

        window.business_report_checks["business_products"].setChecked(False)
        window.preview_report()
        assert "QA Protein Bar" not in window.report_summary.toPlainText()
        assert "QA Calisthenics Plan" in window.report_summary.toPlainText()
        service_only = tmp_path / "service-only.csv"
        window.export_report_csv(service_only, today, report_type="Employee Details", eod_options=window.selected_eod_options())
        with service_only.open(encoding="utf-8-sig", newline="") as stream:
            service_csv = "\n".join(",".join(row) for row in csv.reader(stream))
        assert "QA Calisthenics Plan" in service_csv
        assert "QA Protein Bar" not in service_csv
    finally:
        window.close()
        database.conn.close()
        qt_app.processEvents()


def test_responsive_dashboard_and_tooltip_use_visible_high_contrast_layout(tmp_path, monkeypatch):
    qt_app, database, window = _make_window(tmp_path, monkeypatch)
    try:
        window.resize(1100, 760)
        window._apply_responsive_layout()
        qt_app.processEvents()

        assert window._kpi_column_count in (3, 4)
        assert window._dashboard_column_count in (1, 2)
        assert window.kpis.count() == 5
        assert window.dashboard_grid.count() == 5
        assert window.header_clock_frame.width() == 254
        assert window.att_table.horizontalHeader().stretchLastSection()
        positions = [window.kpis.getItemPosition(index) for index in range(window.kpis.count())]
        assert len({(row, column) for row, column, _row_span, _column_span in positions}) == 5

        tooltip = app.ensure_accessible_tooltip_controller()
        assert tooltip is not None
        style = tooltip.popup.styleSheet()
        assert "background-color: #FFFFFF" in style
        assert "color: #111827" in style
        assert "border: 2px solid #6D28D9" in style
        assert window.brand_logo.pixmap().size() == app.QSize(68, 68)
    finally:
        window.close()
        database.conn.close()
        qt_app.processEvents()


def test_add_task_preserves_selected_employee_and_task_date(tmp_path, monkeypatch):
    qt_app, database, window = _make_window(tmp_path, monkeypatch)
    try:
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        employee_ids = {}
        for name in ("Gopi", "Sai Teja"):
            cursor = database.conn.execute(
                "INSERT INTO employees(name,joining_date,active,created_at) VALUES(?,?,?,?)",
                (name, "", 1, timestamp),
            )
            employee_ids[name] = cursor.lastrowid
        database.conn.commit()
        window.refresh_task_filter()
        window.task_filter.setCurrentIndex(window.task_filter.findData(employee_ids["Sai Teja"]))
        selected_date = app.QDate(2026, 10, 8)
        window.task_date.setDate(selected_date)

        captured = {}
        base_dialog = app.TaskDialog

        class InspectTaskDialog(base_dialog):
            def __init__(self, *args, **kwargs):
                super().__init__(*args, **kwargs)
                captured["employee_id"] = self.employee.currentData()
                captured["dialog"] = self

            def exec(self):
                return app.QDialog.DialogCode.Rejected

        monkeypatch.setattr(app, "TaskDialog", InspectTaskDialog)
        window.add_task()

        assert captured["employee_id"] == employee_ids["Sai Teja"]
        assert captured["dialog"].date.date().toString("yyyy-MM-dd") == "2026-10-08"
    finally:
        window.close()
        database.conn.close()
        qt_app.processEvents()


def test_tracker_add_task_persists_exposed_proof_link(tmp_path, monkeypatch):
    qt_app, database, window = _make_window(tmp_path, monkeypatch)
    try:
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        cursor = database.conn.execute(
            "INSERT INTO employees(name,joining_date,active,created_at) VALUES(?,?,?,?)",
            ("Sai Teja", "", 1, timestamp),
        )
        employee_id = cursor.lastrowid
        database.conn.commit()
        window.refresh_tracker_employees()
        window.tracker_employee.setCurrentIndex(window.tracker_employee.findData(employee_id))

        proof_url = "https://docs.google.com/document/d/tracker-proof/edit?usp=sharing"
        base_dialog = app.TaskDialog

        class ExposedTaskDialog(base_dialog):
            def __init__(self, *args, **kwargs):
                super().__init__(*args, **kwargs)
                self.title.setText("Tracker proof task")
                self.proof.setPlainText(proof_url)
                self.expose_link.setChecked(True)

            def exec(self):
                return app.QDialog.DialogCode.Accepted

        monkeypatch.setattr(app, "TaskDialog", ExposedTaskDialog)
        window.add_tracker_task()

        saved = database.conn.execute(
            "SELECT proof_url, expose_link FROM tasks WHERE employee_id=? AND title=?",
            (employee_id, "Tracker proof task"),
        ).fetchone()
        assert saved["proof_url"] == proof_url
        assert saved["expose_link"] == 1
    finally:
        window.close()
        database.conn.close()
        qt_app.processEvents()


def test_combobox_wheel_rotation_does_not_change_selection(monkeypatch):
    from PySide6.QtCore import QPoint, QPointF, Qt
    from PySide6.QtGui import QWheelEvent
    from PySide6.QtWidgets import QComboBox, QScrollArea, QVBoxLayout, QWidget

    monkeypatch.setenv("QT_QPA_PLATFORM", "offscreen")
    qt_app = app.QApplication.instance() or app.QApplication([])
    app.ensure_combobox_wheel_filter(qt_app)

    scroll = QScrollArea()
    container = QWidget()
    container.setMinimumHeight(2000)
    layout = QVBoxLayout(container)
    combo = QComboBox()
    combo.addItems(["Option A", "Option B", "Option C"])
    combo.setCurrentIndex(0)
    layout.addWidget(combo)
    scroll.setWidget(container)
    scroll.resize(300, 300)
    scroll.show()

    # Rotating mouse wheel over closed combobox should NOT change selection
    wheel_ev = QWheelEvent(
        QPointF(10, 10),
        QPointF(10, 10),
        QPoint(0, 0),
        QPoint(0, -120),
        Qt.MouseButton.NoButton,
        Qt.KeyboardModifier.NoModifier,
        Qt.ScrollPhase.NoScrollPhase,
        False,
    )
    qt_app.sendEvent(combo, wheel_ev)

    assert combo.currentIndex() == 0
    assert combo.currentText() == "Option A"
    assert scroll.verticalScrollBar().value() > 0

    # User explicitly selecting an option changes the value normally
    combo.setCurrentIndex(1)
    assert combo.currentIndex() == 1
    assert combo.currentText() == "Option B"




def test_print_center_options_match_task_manager_columns(tmp_path, monkeypatch):
    from docx import Document
    import reportlab.platypus
    import json

    qt_app, database, window = _make_window(tmp_path, monkeypatch)
    try:
        today = date.today().isoformat()
        now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        # 1. Assert all Task Manager fields are present in eod_checks
        expected_keys = [
            "target_qty", "actual_qty", "completion_pct", "priority",
            "custom_values", "next_action", "tasks", "details", "due_time",
            "proof", "proof_notes", "expose_link", "remarks"
        ]
        for key in expected_keys:
            assert key in window.eod_checks, f"Missing {key} in window.eod_checks"

        # 2. Insert test employee and task with full Task Manager fields
        emp_id = database.conn.execute(
            "INSERT INTO employees(name,employee_code,department,designation,active,created_at) VALUES(?,?,?,?,?,?)",
            ("Print QA Athlete", "PQA01", "Calisthenics", "Senior Coach", 1, now_str),
        ).lastrowid

        custom_json = json.dumps([{"key": "Tracked By", "value": "Coach John"}, {"key": "Equipment", "value": "Rings"}])
        database.conn.execute(
            """INSERT INTO tasks(employee_id, task_date, title, department, details, due_time, priority, status, progress,
                                proof_url, proof_note, manager_note, expose_link, target_qty, actual_qty, next_action, custom_values,
                                created_at, updated_at)
               VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (emp_id, today, "Muscle-up sets", "Strength", "Strict form", "11:00", "High", "In progress", 75,
             "https://example.com/mu", "Great angle", "Manager reviewing", 0, "50 reps", "38 reps", "Review grip video",
             custom_json, now_str, now_str)
        )
        database.conn.commit()

        # 3. Test with all Task Manager options checked
        opts_all = window.selected_eod_options()
        assert opts_all["target_qty"] is True
        assert opts_all["actual_qty"] is True
        assert opts_all["completion_pct"] is True
        assert opts_all["priority"] is True
        assert opts_all["custom_values"] is True
        assert opts_all["next_action"] is True

        # Preview text
        preview = window.report_preview_text(today, emp_id)
        assert "Target Qty: 50 reps" in preview
        assert "Actual Qty: 38 reps" in preview
        assert "Completion %: 75%" in preview
        assert "Priority: High" in preview
        assert "Custom Values: Tracked By: Coach John  |  Equipment: Rings" in preview
        assert "Next Action: Review grip video" in preview

        # Word document
        docx_path = tmp_path / "test_tm_options.docx"
        window.make_docx(docx_path, today, emp_id, "Audit / EOD", opts_all)
        doc = Document(docx_path)
        doc_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Target Qty: 50 reps" in doc_text
        assert "Actual Qty: 38 reps" in doc_text
        assert "Completion %: 75%" in doc_text
        assert "Priority: High" in doc_text
        assert "Custom Values: Tracked By: Coach John  |  Equipment: Rings" in doc_text
        assert "Next Action: Review grip video" in doc_text

        # CSV export
        csv_path = tmp_path / "test_tm_options.csv"
        window.export_report_csv(csv_path, today, emp_id, "Audit / EOD", opts_all)
        with open(csv_path, encoding="utf-8-sig") as f:
            csv_rows = list(csv.reader(f))
        headers = csv_rows[0]
        assert "Target Qty" in headers
        assert "Actual Qty" in headers
        assert "Completion %" in headers
        assert "Priority" in headers
        assert "Custom Values" in headers
        assert "Next Action" in headers
        data_row = csv_rows[1]
        assert "50 reps" in data_row
        assert "38 reps" in data_row
        assert "75%" in data_row
        assert "High" in data_row
        assert "Tracked By: Coach John  |  Equipment: Rings" in data_row
        assert "Review grip video" in data_row

        # PDF export
        captured_story = []
        class CaptureDoc:
            def __init__(self, *args, **kwargs): pass
            def build(self, story, **kwargs): captured_story.extend(story)
        monkeypatch.setattr(reportlab.platypus, "SimpleDocTemplate", CaptureDoc)
        pdf_path = tmp_path / "test_tm_options.pdf"
        window.make_pdf(pdf_path, today, emp_id, "Audit / EOD", opts_all)
        pdf_text = " ".join(getattr(f, "text", "") for f in captured_story if hasattr(f, "text"))
        assert "Target Qty: 50 reps" in pdf_text
        assert "Actual Qty: 38 reps" in pdf_text
        assert "Completion %: 75%" in pdf_text
        assert "Priority: High" in pdf_text
        assert "Custom Values:" in pdf_text and "Coach John" in pdf_text and "Rings" in pdf_text
        assert "Next Action: Review grip video" in pdf_text

        # 4. Now uncheck Task Manager options and verify exclusion
        window.eod_checks["target_qty"].setChecked(False)
        window.eod_checks["actual_qty"].setChecked(False)
        window.eod_checks["completion_pct"].setChecked(False)
        window.eod_checks["priority"].setChecked(False)
        window.eod_checks["custom_values"].setChecked(False)
        window.eod_checks["next_action"].setChecked(False)
        opts_unchecked = window.selected_eod_options()

        preview_un = window.report_preview_text(today, emp_id)
        assert "Target Qty" not in preview_un
        assert "Actual Qty" not in preview_un
        assert "Completion %" not in preview_un
        assert "Priority: High" not in preview_un
        assert "Custom Values" not in preview_un
        assert "Next Action" not in preview_un

        docx_un_path = tmp_path / "test_tm_options_un.docx"
        window.make_docx(docx_un_path, today, emp_id, "Audit / EOD", opts_unchecked)
        doc_un = Document(docx_un_path)
        doc_un_text = "\n".join(p.text for p in doc_un.paragraphs)
        assert "Target Qty" not in doc_un_text
        assert "Actual Qty" not in doc_un_text
        assert "Completion %" not in doc_un_text
        assert "Priority: High" not in doc_un_text
        assert "Custom Values" not in doc_un_text
        assert "Next Action" not in doc_un_text

        csv_un_path = tmp_path / "test_tm_options_un.csv"
        window.export_report_csv(csv_un_path, today, emp_id, "Audit / EOD", opts_unchecked)
        with open(csv_un_path, encoding="utf-8-sig") as f:
            csv_un_headers = list(csv.reader(f))[0]
        assert "Target Qty" not in csv_un_headers
        assert "Actual Qty" not in csv_un_headers
        assert "Completion %" not in csv_un_headers
        assert "Priority" not in csv_un_headers
        assert "Custom Values" not in csv_un_headers
        assert "Next Action" not in csv_un_headers
    finally:
        window.close()
        database.conn.close()
        qt_app.processEvents()
